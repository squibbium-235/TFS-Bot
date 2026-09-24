"""Regenerate the Sanctuary Servo server-owner manual.

python-docx is a one-off tool for this script. It is not a bot
dependency and is not listed in requirements.txt.

    pip install python-docx
    python docs/build_server_owner_manual.py

The Word file is written next to this script.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT_PATH = Path(__file__).with_name(
    "Sanctuary-Servo-Server-Owner-Manual.docx"
)

INK = RGBColor(0x1C, 0x16, 0x0E)
BRASS = RGBColor(0x6B, 0x53, 0x2A)
BODY = RGBColor(0x2A, 0x26, 0x22)


def shade_cell(cell, fill: str) -> None:
    """Set a table cell background. fill is a hex colour without '#'."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    tc_pr.append(shading)


def set_run_font(run, *, name: str, size: int, bold: bool = False, colour: RGBColor = BODY) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = colour


def add_paragraph(document: Document, text: str, *, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_font(run, name="Calibri", size=11)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.space_before = Pt(0)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_run_font(run, name="Calibri", size=11)
        paragraph.paragraph_format.space_after = Pt(2)


def add_numbers(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        run = paragraph.add_run(item)
        set_run_font(run, name="Calibri", size=11)
        paragraph.paragraph_format.space_after = Pt(2)


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_heading(text, level=level)
    size = {1: 18, 2: 14, 3: 12}[level]
    for run in paragraph.runs:
        set_run_font(
            run,
            name="Cambria",
            size=size,
            bold=True,
            colour=BRASS if level > 1 else INK,
        )
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(6)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(header)
        set_run_font(run, name="Calibri", size=10, bold=True, colour=RGBColor(0xFF, 0xF8, 0xF0))
        shade_cell(cell, "3D3428")

    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row):
            cell = table.rows[row_index].cells[column_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(value)
            set_run_font(run, name="Calibri", size=10)
            if row_index % 2 == 0:
                shade_cell(cell, "F6F1E8")

    document.add_paragraph()


def add_note(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    label = paragraph.add_run("Note. ")
    set_run_font(label, name="Calibri", size=11, bold=True, colour=BRASS)
    body = paragraph.add_run(text)
    set_run_font(body, name="Calibri", size=11)
    paragraph.paragraph_format.space_after = Pt(8)


def build() -> Document:
    document = Document()

    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    kicker = title.add_run("SANCTUARY SERVO")
    set_run_font(kicker, name="Cambria", size=12, bold=True, colour=BRASS)

    heading = document.add_paragraph()
    heading_run = heading.add_run("Server owner manual")
    set_run_font(heading_run, name="Cambria", size=26, bold=True, colour=INK)
    heading.paragraph_format.space_after = Pt(4)

    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run(
        "For the person who installs, configures, and looks after the bot on a Discord server."
    )
    set_run_font(subtitle_run, name="Calibri", size=12, colour=BRASS)
    subtitle.paragraph_format.space_after = Pt(10)

    add_paragraph(
        document,
        "This manual describes the bot as it behaves in this repository. "
        "It is written for a server owner, not for someone changing the Python code. "
        "Log lines still use the name TFSBot. That is the program name. The product name is Sanctuary Servo.",
    )

    add_heading(document, "Contents", 1)
    add_bullets(
        document,
        [
            "1. What the bot does",
            "2. Discord application, intents, and invite",
            "3. Permissions the bot needs in the server",
            "4. Install, configure, and run",
            "5. Permission levels",
            "6. Web UI",
            "7. Verification",
            "8. Forms",
            "9. Welcome messages",
            "10. Custom commands",
            "11. DM templates, embeds, and uploads",
            "12. Moderation profiles and diagnostics",
            "13. Backups",
            "14. Common operator tasks",
            "15. Troubleshooting",
        ],
    )

    add_heading(document, "1. What the bot does", 1)
    add_paragraph(
        document,
        "Sanctuary Servo is a Discord bot for running a server where new people apply before they are fully let in. "
        "Staff review those applications in Discord. The same bot stores the configuration for that workflow.",
    )
    add_paragraph(document, "In day-to-day use it provides:")
    add_bullets(
        document,
        [
            "A verification panel. People press Verify, fill in a form, and staff approve, reject, kick, ban, or ask further questions.",
            "Forms you can edit, including the verification form and other forms you publish as their own panels.",
            "Direct messages for each verification outcome, with wording you can change.",
            "A welcome message sent after an application is approved. It is not sent merely because someone joined.",
            "Custom prefix commands, built from a list of actions.",
            "Permission levels that decide who may use staff commands, and who may open the Web UI.",
            "An optional Web UI for owners. Viewers can open the overview only.",
            "An encrypted database, and encrypted backup files.",
        ],
    )
    add_paragraph(
        document,
        "Two small commands are always available: /ping and !ping check that the bot is responding, "
        "and /info and !info show a short status embed. The default prefix is !.",
    )

    add_heading(document, "2. Discord application, intents, and invite", 1)
    add_paragraph(
        document,
        "Create the application in the Discord Developer Portal, add a bot user, and copy the bot token into .env as DISCORD_TOKEN. "
        "Do not paste the token into a channel, a backup you share, or this manual.",
    )
    add_heading(document, "Privileged intents", 2)
    add_paragraph(
        document,
        "Turn both of these on for the bot in the Developer Portal. The process asks Discord for them at startup. "
        "If they are off, Discord will refuse the connection or the related features will not see the data they need.",
    )
    add_table(
        document,
        ["Intent", "Why it is required"],
        [
            [
                "Message Content Intent",
                "Prefix commands (!ping, !info, and custom commands) need to read message text.",
            ],
            [
                "Server Members Intent",
                "Verification, welcome messages, and invite tracking need the member list and member join and leave events.",
            ],
        ],
    )
    add_heading(document, "Invite", 2)
    add_paragraph(document, "Invite the bot with both of these scopes:")
    add_bullets(
        document,
        [
            "bot",
            "applications.commands",
        ],
    )
    add_paragraph(
        document,
        "Without applications.commands, slash commands will not appear. "
        "A global sync can take a while to show up in Discord. "
        "Set TEST_GUILD_ID to one server's ID while you are setting the bot up if you want that server to receive commands immediately. "
        "Clear it when you want commands published globally. A guild sync only publishes commands to that one server.",
    )

    add_heading(document, "3. Permissions the bot needs in the server", 1)
    add_paragraph(
        document,
        "Put the bot's role above any role it must assign or remove. Discord will not let a bot manage a role that sits level with it or above it. "
        "/diagnostics reports that case as an error.",
    )
    add_paragraph(document, "Server-wide permissions used by the built-in checks and actions:")
    add_table(
        document,
        ["Permission", "Used for"],
        [
            ["Manage Roles", "Giving and removing the approval role, and custom-command role actions."],
            ["Kick Members", "Rejecting an application with a kick."],
            ["Ban Members", "Rejecting an application with a ban."],
            ["Manage Server", "Reading the invite list. Without it, invite tracking stays unsynchronised."],
            ["Moderate Members", "The person running /modprofile needs this. It is a permission on the staff member, not only on the bot."],
        ],
    )
    add_paragraph(document, "In the review channel, the log channel, and the welcome channel, the bot needs:")
    add_bullets(
        document,
        [
            "View Channel",
            "Send Messages",
            "Embed Links",
            "Attach Files",
            "Read Message History",
            "Create Public Threads, and Send Messages in Threads, so staff can question an applicant",
            "Manage Threads, so a questioning thread can be locked and archived when the case closes",
        ],
    )
    add_paragraph(
        document,
        "The bot also sends direct messages. If a person has DMs closed, approval, rejection, and questioning messages to them will fail even when the rest of the setup is correct.",
    )

    add_heading(document, "4. Install, configure, and run", 1)
    add_heading(document, "Files that matter", 2)
    add_bullets(
        document,
        [
            ".env holds the token, the database key, and Web UI settings. It is not committed. Start from .env.example.",
            "data/tfsbot.sqlite3 is the default encrypted database. Every store shares this file.",
            "data/uploads/ holds images used by the embed builder and welcome messages.",
            "data/forms/verification.json is the fallback verification form if the database does not have one yet.",
            "data/restore_safety/ is created when you restore a backup. It is a copy of what was replaced.",
        ],
    )
    add_heading(document, "Required settings", 2)
    add_paragraph(
        document,
        "From the repository root, create a virtual environment, install requirements.txt, copy .env.example to .env, and fill it in. Then run:",
    )
    add_paragraph(document, "python -m src.main")
    add_paragraph(document, "The process will not stay up if these are missing or invalid:")
    add_table(
        document,
        ["Setting", "What to put"],
        [
            ["DISCORD_TOKEN", "The bot token."],
            [
                "TFSBOT_DATABASE_KEY",
                "Exactly 64 hexadecimal characters. Generate one with python -c \"import secrets; print(secrets.token_hex(32))\". Keep a copy somewhere other than the server.",
            ],
        ],
    )
    add_note(
        document,
        "If you lose the database key, the database cannot be opened. A backup does not change the key. "
        "The file inside the backup is still encrypted with the same key. Include .env in a backup only when you accept that the file then contains the token and the key.",
    )
    add_heading(document, "Other settings", 2)
    add_table(
        document,
        ["Setting", "Default", "Purpose"],
        [
            ["BOT_PREFIX", "!", "Prefix for !ping, !info, and custom commands."],
            ["TEST_GUILD_ID", "empty", "Sync slash commands to this server only."],
            ["APPLICATION_DB_PATH", "data/tfsbot.sqlite3", "Where the encrypted database is stored."],
            ["BOT_DEV_USER_IDS", "empty", "Comma-separated user IDs that are always treated as owner."],
            ["WEBUI_ENABLED", "false", "Start the Web UI with the bot."],
            ["WEBUI_HOST", "127.0.0.1", "Bind address. 127.0.0.1 means this machine only."],
            ["WEBUI_PORT", "5050", "Web UI port."],
            ["WEBUI_PASSWORD_LOGIN_ENABLED", "true", "Allow the username and password form."],
            ["WEBUI_USER_1_USERNAME / PASSWORD", "empty", "First password account. Both fields are required if either is set."],
            ["WEBUI_USER_2_USERNAME / PASSWORD", "empty", "Second password account. There is no third."],
            ["WEBUI_AUTH_DISCORD_ENABLED", "false", "Allow Log in with Discord."],
            ["DISCORD_OAUTH_CLIENT_ID", "empty", "Required when Discord login is on."],
            ["DISCORD_OAUTH_CLIENT_SECRET", "empty", "Required when Discord login is on."],
            ["DISCORD_OAUTH_REDIRECT_URI", "empty", "Must match the Developer Portal exactly. Usually ends in /auth/discord/callback."],
            ["WEBUI_DISCORD_GUILD_ID", "empty", "The server whose roles are checked at Discord login."],
            ["WEBUI_DISCORD_ALLOWED_ROLE_IDS", "empty", "Roles that may try to log in, until the Permissions page stores its own list."],
            ["WEBUI_DISCORD_OWNER_ROLE_IDS", "empty", "Roles that become Web UI owners."],
            ["WEBUI_DISCORD_VIEWER_ROLE_IDS", "empty", "Roles that become Web UI viewers."],
        ],
    )
    add_paragraph(
        document,
        "Boolean settings accept 1, true, yes, y, or on. Anything else that is set, including an empty value, is false. "
        "If the Web UI is enabled and Discord login is off, password login must be on and at least one account must be complete. "
        "If Discord login is on, the client id, client secret, redirect URI, and guild id are all required.",
    )
    add_heading(document, "Docker", 2)
    add_paragraph(
        document,
        "compose.yaml builds the image and mounts ./data onto /app/data so the database survives a new container. "
        "It reads .env. The Web UI port is exposed to other Compose services only. It is not published on the host.",
    )
    add_paragraph(
        document,
        "Inside the container, the default WEBUI_HOST of 127.0.0.1 listens on the container's own loopback, which you cannot reach from outside. "
        "To use the Web UI from Docker, set WEBUI_HOST=0.0.0.0 and publish the port yourself. "
        "Treat that port as an admin console. Do not leave it open to the public internet.",
    )

    add_heading(document, "5. Permission levels", 1)
    add_paragraph(
        document,
        "Command access uses four levels, from lowest to highest: public, staff, admin, and owner. "
        "A higher level can use commands that require a lower level. These levels are separate from Discord's own role permissions, except where a command also asks for one, such as Moderate Members on /modprofile.",
    )
    add_paragraph(document, "A person's level is decided in this order:")
    add_numbers(
        document,
        [
            "A user ID listed in BOT_DEV_USER_IDS is owner.",
            "The Discord server owner is owner.",
            "A member with the configured owner role is owner. The admin role is admin. The staff role is staff.",
            "If none of those roles match, Administrator or Manage Server counts as admin.",
            "Everyone else is public.",
        ],
    )
    add_note(
        document,
        "A staff role stays staff even when that person also has Administrator. "
        "The role match is checked first, so the Discord permission is only a fallback.",
    )
    add_paragraph(
        document,
        "Set the three roles with /permissions set-role, or on the Web UI Permissions page. "
        "/permissions my-level shows your own level and is public. "
        "The other /permissions commands default to owner: view, set-role, clear-role, set-command, and reset-command.",
    )
    add_paragraph(document, "Built-in command defaults:")
    add_table(
        document,
        ["Commands", "Default level"],
        [
            ["/ping, /info", "Public"],
            ["/permissions my-level", "Public"],
            ["/form preview, /form submissions", "Staff"],
            ["/form (other subcommands), /verification, /welcome, /custom-command, /diagnostics, /permissions (other subcommands)", "Owner"],
            ["/modprofile", "Public in this bot's table, plus Discord's Moderate Members permission"],
        ],
    )
    add_paragraph(
        document,
        "You can raise or lower a command with /permissions set-command. "
        "The command key uses dots and underscores: /verification review-channel is verification.review_channel. "
        "A child command uses its own override, then its parent's override, then the built-in default. "
        "/permissions reset-command removes the override.",
    )
    add_paragraph(
        document,
        "Slash commands that fail the level check reply privately with the required level and the person's level. "
        "A direct message has no roles to compare, so the level check is skipped there. "
        "If the permission store is missing, the command is blocked.",
    )

    add_heading(document, "6. Web UI", 1)
    add_paragraph(
        document,
        "The Web UI starts in the same process as the bot, only when WEBUI_ENABLED is true. "
        "Open http://127.0.0.1:5050/ if you kept the defaults and you are on the same machine. "
        "The log line says Starting web UI when it launches.",
    )
    add_heading(document, "Signing in", 2)
    add_bullets(
        document,
        [
            "Password login uses WEBUI_USER_1 or WEBUI_USER_2. Password accounts are owners.",
            "Discord login sends the person to Discord and back to /auth/discord/callback. The redirect URI in .env and in the Developer Portal must be identical.",
            "Discord login only succeeds for a member of WEBUI_DISCORD_GUILD_ID who matches an allowed role.",
            "Owner roles win over viewer roles. A viewer can open the overview. Every other page shows Access denied.",
            "Role lists saved on the Permissions page replace the environment lists for that server and level. Clearing the saved list falls back to the environment.",
            "If both methods are on, Discord is the main button and the password form is labelled as an emergency log in.",
        ],
    )
    add_heading(document, "Sessions", 2)
    add_bullets(
        document,
        [
            "A signed-in session lasts eight hours from the moment of login, and also ends after one hour without a request.",
            "Moving around the site refreshes the idle timer. It does not extend the eight-hour limit.",
            "The cookie is HTTP-only and SameSite=Lax.",
            "Unsafe requests (the forms you submit) must include the session CSRF token. A missing or stale token returns HTTP 400. Reload the page, or sign in again, and retry.",
            "Uploads are refused above 10 MiB.",
        ],
    )
    add_heading(document, "Pages", 2)
    add_paragraph(document, "The navigation is hidden for viewers except Overview and Log out.")
    add_table(
        document,
        ["Page", "Path", "What you do there"],
        [
            ["Overview", "/", "Queue counts, setup health, and recent applications. Any signed-in role."],
            ["Embed Builder", "/embed-builder", "Compose an embed, save it, and post it to a channel."],
            ["Custom Commands", "/custom-commands", "Create prefix commands and edit their actions."],
            ["DM Templates", "/dm-templates", "Wording sent to applicants for each verification outcome."],
            ["Forms", "/forms", "Create and edit modal forms. Submissions open at /forms/view."],
            ["Uploads", "/uploads-manager", "Image files used by embeds. PNG, JPEG, GIF, and WebP."],
            ["Verification", "/verification", "Channels, roles, automod, the panel, and invite status. Welcome is /verification/welcome."],
            ["Permissions", "/permissions", "Web UI roles and slash-command levels for one server."],
            ["Backups", "/backups", "Create or restore an encrypted .tfsbackup file."],
        ],
    )
    add_paragraph(
        document,
        "Most pages start with a server menu, because the bot can be in more than one server. "
        "Changing the menu reloads that server's settings. Successful changes are written to the audit log.",
    )

    add_heading(document, "7. Verification", 1)
    add_paragraph(
        document,
        "Verification is the main workflow. A public panel asks people to press Verify and complete the form. "
        "The submission is posted in the review channel with Approve, Reject, Kick, Ban, and Question. "
        "The result is also written to the log channel when one is set.",
    )
    add_heading(document, "Set it up", 2)
    add_numbers(
        document,
        [
            "Choose the review channel. This is where staff work the queue. /verification review-channel, or the Verification page.",
            "Choose the log channel if you want a separate record. /verification log-channel.",
            "Set the role to give on approval, and optionally a role to remove (for example an Unverified role). The bot's role must be above both.",
            "Check the form. The default key is verification. Edit it under Forms, then post the panel again so new applicants see the new questions. People who already have the form open are not rewritten.",
            "Post the panel with /verification panel or the Web UI. You can attach a large image and a thumbnail. If you do not upload a thumbnail, the server icon is used.",
            "Run /diagnostics and clear anything marked as an error.",
        ],
    )
    add_paragraph(
        document,
        "/verification setup opens an interactive configuration panel in Discord. "
        "/verification status shows the current channels, roles, automod, invite tracking, and welcome state. "
        "/verification help lists the subcommands. Administration replies are visible only to the person who ran the command.",
    )
    add_heading(document, "Staff actions", 2)
    add_table(
        document,
        ["Action", "What happens"],
        [
            ["Approve", "Applicant is approved, the add-role is given, the remove-role is taken off, and the approved DM is sent. A welcome message is queued if welcome is enabled."],
            ["Reject", "Application is denied and the rejected DM is sent. The person stays in the server."],
            ["Kick", "Application is denied, the person is kicked, and the kicked DM is sent."],
            ["Ban", "Application is denied, the person is banned, and the banned DM is sent."],
            ["Question", "A thread is opened on the review message. Staff messages in that thread are forwarded to the applicant's DMs, and their replies come back to the thread."],
        ],
    )
    add_paragraph(
        document,
        "A message in the questioning thread that starts with // stays in the thread. It is not forwarded and it is not treated as a command. "
        "Example: // this answer seems suspicious. "
        "When the case closes, the bot tries to lock and archive the thread. That needs Manage Threads.",
    )
    add_paragraph(
        document,
        "If the applicant leaves while an application is still pending, the bot records that departure. "
        "A kick or ban started by staff is not also recorded as the applicant walking away.",
    )
    add_heading(document, "Automod", 2)
    add_paragraph(
        document,
        "Verification automod is a list of blocked terms checked against applications. "
        "Turn it on or off with /verification automod-enabled. Add, remove, and list terms with automod-add, automod-remove, and automod-list, or edit the list on the Verification page.",
    )
    add_paragraph(
        document,
        "On the Web UI, one term per line. Blank lines are dropped. A line whose first character is # is a comment and is not stored as a term. "
        "You can insert the built-in default list from the page. That happens after the text box is saved, so both are kept.",
    )
    add_heading(document, "Cancelling applications", 2)
    add_bullets(
        document,
        [
            "/verification cancel-user cancels one pending application.",
            "/verification cancel-all cancels every pending application in that server.",
            "The Web UI asks you to type CANCEL before it will cancel.",
        ],
    )
    add_heading(document, "Invites", 2)
    add_paragraph(
        document,
        "On startup the bot reads each server's invites once, then keeps that cache up to date. "
        "A reconnect does not rebuild it from scratch. Fetching invites needs Manage Server. "
        "If Discord refuses, tracking stays unsynchronised and /diagnostics shows a warning rather than an error. "
        "The tracked inviter can appear in the application and in welcome messages as {inviter}.",
    )

    add_heading(document, "8. Forms", 1)
    add_paragraph(
        document,
        "Forms are Discord modals: a short title and text questions. "
        "Discord allows at most five questions on one modal and a title of at most 45 characters, so a longer form is split across pages. "
        "Question labels are limited to 45 characters. A placeholder is limited to 100.",
    )
    add_paragraph(
        document,
        "A form key is 1 to 40 characters: lower-case letters, digits, and underscores. "
        "A question key is 1 to 80 characters in the same style. "
        "The verification form uses the key verification unless you point the server at another key.",
    )
    add_paragraph(document, "You can edit forms in the Web UI or with /form:")
    add_table(
        document,
        ["Command", "Level", "Purpose"],
        [
            ["/form create", "Owner", "Create a form."],
            ["/form list, /form view", "Owner", "See forms and their questions."],
            ["/form add, /form edit, /form delete, /form move", "Owner", "Change questions. Styles are short answer or paragraph."],
            ["/form delete-form", "Owner", "Delete a whole form. The form currently used for verification cannot be deleted."],
            ["/form reset-verification", "Owner", "Put the verification form back to the built-in default."],
            ["/form publish", "Owner", "Post a panel for a form that is not the verification form."],
            ["/form preview", "Staff", "Preview the modal pages."],
            ["/form submissions", "Staff", "Recent submissions. The Web UI viewer is /forms/view."],
        ],
    )
    add_note(
        document,
        "Changing the verification form affects forms opened after the change. "
        "Post the verification panel again after you edit it. "
        "The Web UI does not keep placeholder text or length limits when you save a question from that page.",
    )

    add_heading(document, "9. Welcome messages", 1)
    add_paragraph(
        document,
        "A welcome message is an embed posted in a channel you choose, after a verification application is approved. "
        "Joining the server does not send it. "
        "The bot checks for new approvals about once a second. Only approvals made at or after the moment you enabled welcome are eligible, and each application is welcomed at most once.",
    )
    add_paragraph(
        document,
        "Configure it with /welcome or on /verification/welcome. "
        "Set a channel before you enable it. Enable refuses to turn the feature on when the channel is missing. "
        "Disable only switches the feature off; it does not wipe the embed. "
        "Preview builds the embed and does not post it. Test posts a sample, including from the unsaved Web UI form.",
    )
    add_paragraph(document, "Placeholders:")
    add_table(
        document,
        ["Placeholder", "Replaced with"],
        [
            ["{user}", "A mention of the approved member"],
            ["{username}", "Their Discord username"],
            ["{display_name}", "Their name in the server"],
            ["{user_id}", "Their user ID"],
            ["{server}", "The server name"],
            ["{member_count}", "The current member count"],
            ["{inviter}", "The tracked inviter, when one was recorded"],
            ["{avatar}", "Their avatar URL"],
        ],
    )
    add_paragraph(
        document,
        "The default embed is titled Welcome {user}! and uses Discord blurple, #5865F2, until you choose another colour. "
        "A colour is six hex digits, with an optional # or 0x in front. "
        "An author icon is kept only when an author name is set. "
        "On the welcome page, a new upload is used first, then a URL you type, then an image already stored. "
        "A URL that contains { is left as text in the preview, because a placeholder such as {avatar} is not an image the browser can fetch.",
    )
    add_paragraph(
        document,
        "Slash subcommands, all owner by default: setup, edit, extras, channel, add-field, fields, remove-field, clear-fields, enable, disable, status, preview, and test.",
    )

    add_heading(document, "10. Custom commands", 1)
    add_paragraph(
        document,
        "Custom commands are prefix commands, not slash commands. "
        "With the default prefix, a command named rules runs when someone sends !rules. "
        "Names are 1 to 32 characters: lower-case letters, digits, hyphens, and underscores. "
        "New commands start enabled and with no actions.",
    )
    add_paragraph(
        document,
        "Create and edit them with /custom-command (owner by default) or on the Custom Commands page. "
        "Each command has a description, an enabled flag, a required level (public, staff, admin, or owner), a cooldown from 0 to 86400 seconds, and whether the triggering message should be deleted. "
        "The cooldown is remembered in memory only. A restart clears it. The cooldown is used up even if a later action fails.",
    )
    add_paragraph(document, "Actions run in order. You can add:")
    add_bullets(
        document,
        [
            "Send a message.",
            "Send an embed, then add fields to that embed. An embed can hold at most 25 fields.",
            "Add a reaction, either to the triggering message or to the latest bot response.",
            "Add or remove a role, either on the person who ran the command or on a mentioned person.",
            "Delete the triggering message or the latest bot response.",
        ],
    )
    add_paragraph(
        document,
        "Useful subcommands: create, delete, list, view, settings, add-message, add-embed, add-embed-field, add-reaction, add-role, remove-role, add-delete, remove-action, move-action, clear-actions, and placeholders.",
    )
    add_paragraph(document, "Placeholders in message and embed text:")
    add_table(
        document,
        ["Placeholder", "Meaning"],
        [
            ["{user}, {display_name}, {mention}, {user_id}", "The person who ran the command"],
            ["{target}, {target_mention}, {target_id}", "The mentioned person, when the action uses one"],
            ["{server}, {server_id}, {member_count}", "This server"],
            ["{channel}, {channel_id}", "The channel the command was used in"],
            ["{args}", "Everything after the command name"],
            ["{arg1}, {arg2}, {arg3}…", "Individual arguments. A missing one becomes empty text."],
            ["{command}, {prefix}", "The command name, and the bot prefix"],
            ["{newline}", "A line break"],
            ["{random:one|two|three}", "One of the options, chosen after the other placeholders are filled in"],
        ],
    )
    add_paragraph(document, "A custom command does not run when:")
    add_bullets(
        document,
        [
            "The message is from a bot, including Sanctuary Servo itself.",
            "It is a direct message.",
            "Discord already treats it as a built-in prefix command. !ping and !info win over a custom command of the same name.",
            "It is posted in an open verification questioning thread. That thread is reserved for the question bridge.",
            "The command is disabled, the person is below its required level, or it is still on cooldown for them.",
        ],
    )

    add_heading(document, "11. DM templates, embeds, and uploads", 1)
    add_heading(document, "DM templates", 2)
    add_paragraph(
        document,
        "Five messages can be edited. Saving text that matches the built-in default clears the saved copy and the default is used again. "
        "An unknown placeholder is left as written, including the braces. If a template cannot be rendered, the raw text is sent so the moderation action still completes.",
    )
    add_table(
        document,
        ["Template", "When it is sent"],
        [
            ["Approved application", "The application was approved."],
            ["Rejected application", "The application was rejected and the person was left in the server."],
            ["Kicked after rejection", "The application was rejected and the person was kicked."],
            ["Banned after rejection", "The application was rejected and the person was banned."],
            ["Questioning opened", "Staff opened a questioning thread. The default asks them to reply to the DM."],
        ],
    )
    add_paragraph(document, "Placeholders: {user}, {user_name}, {user_id}, {server_name}, {moderator}, {moderator_name}, {moderator_id}, {application_id}, {reason}, and {reason_block}.")
    add_paragraph(
        document,
        "{reason} is the reason text, or nothing. "
        "{reason_block} is a blank line, the word Reason, and the text, or nothing when there is no reason, so you can leave it in the template.",
    )
    add_heading(document, "Embed builder and uploads", 2)
    add_paragraph(
        document,
        "The Embed Builder edits a Discord embed and can post it to a channel. Saved embeds stay in the database. "
        "The preview on the right is painted in Discord's colours so you can see the post, not the admin page, before you send it.",
    )
    add_paragraph(
        document,
        "Images live under Uploads. Put author icons, thumbnails, and other pictures in folders so they stay findable. "
        "Accepted types are PNG, JPEG, GIF, and WebP, up to 10 MiB. "
        "A folder can be deleted only when it is empty. "
        "Welcome messages and verification panels can use these files as well.",
    )

    add_heading(document, "12. Moderation profiles and diagnostics", 1)
    add_heading(document, "Moderation profiles", 2)
    add_paragraph(
        document,
        "/modprofile opens or creates a staff profile for one person in the configured log channel. "
        "You can pass a user ID, a username, a display name, or a mention. "
        "The person running the command needs Discord's Moderate Members permission. "
        "In Sanctuary Servo's own permission table the command is not listed, so it stays public unless you set a level with /permissions set-command. "
        "The profile embed includes account and join dates, verification attempts, the Discord moderation log, and staff notes. "
        "If the saved message can no longer be edited, the bot deletes the stored profile and posts a new message and thread.",
    )
    add_heading(document, "Diagnostics", 2)
    add_paragraph(
        document,
        "/diagnostics is an owner command. It replies privately with one field per check. "
        "The embed is red when any check is an error, orange when there are only warnings, and green when everything is healthy. "
        "Warnings do not by themselves make the report unhealthy.",
    )
    add_paragraph(document, "It checks:")
    add_bullets(
        document,
        [
            "Whether the bot is connected.",
            "The review channel and the log channel, including View Channel, Send Messages, Embed Links, Attach Files, and Read Message History.",
            "The approval add-role and remove-role, including whether each role still exists and sits below the bot.",
            "Whether the verification form loads.",
            "Whether the database file exists and can be written. This check does not decrypt the file, so a wrong key can still look healthy here.",
            "Whether the invite cache has been synchronised.",
            "Manage Roles, Kick Members, and Ban Members on the bot.",
        ],
    )
    add_paragraph(
        document,
        "The overview page shows a related summary for the selected server. Use it when you do not want to open Discord.",
    )

    add_heading(document, "13. Backups", 1)
    add_paragraph(
        document,
        "Backups are created from the Web UI Backups page. The file extension is .tfsbackup. "
        "The download name looks like Sanctuary_Servo_Backup_2026-09-24_093000.tfsbackup, in UTC.",
    )
    add_paragraph(document, "A backup contains:")
    add_table(
        document,
        ["Path inside the archive", "Included"],
        [
            ["data/tfsbot.sqlite3", "Always. This is a copy of the encrypted database file, not a new export. The key does not change."],
            ["data/uploads/", "When the folder exists."],
            [".env", "Only if you tick the option. It then contains the token and the database key."],
        ],
    )
    add_paragraph(
        document,
        "The backup password must be at least 10 characters, and the two password fields must match. "
        "The password is not stored. Without it the file cannot be decrypted. "
        "Keep the password with the same care as the database key, but do not write it on the backup filename.",
    )
    add_heading(document, "Restore", 2)
    add_numbers(
        document,
        [
            "Sign out and sign back in. Restore only works for 10 minutes after login. Using the site does not extend that window.",
            "Choose a file whose name ends in .tfsbackup.",
            "Enter the backup password.",
            "Type RESTORE in the confirmation box.",
            "Choose whether to restore uploads and whether to restore .env. Restoring .env requires that the archive actually contains one.",
            "Restart the bot as soon as the page says the restore finished. Do not keep using the Web UI against the replaced database.",
        ],
    )
    add_paragraph(
        document,
        "Before it replaces anything, the current database is copied to data/restore_safety/ followed by a UTC timestamp. "
        "Uploads, when restored, replace the existing uploads folder. "
        "If a later step fails, the database may already have been replaced. Use the safety copy if you need to go back, and restart afterwards.",
    )
    add_note(
        document,
        "Take a backup before you edit forms, permission levels, or welcome text you would be unhappy to lose, and before you move the bot to another machine. "
        "On the new machine you still need the same TFSBOT_DATABASE_KEY unless the backup included .env and you chose to restore it.",
    )

    add_heading(document, "14. Common operator tasks", 1)
    add_heading(document, "A new server", 2)
    add_numbers(
        document,
        [
            "Invite the bot with the bot and applications.commands scopes, and with the permissions in section 3.",
            "Confirm both privileged intents are on, then start the process and watch for Logged in as in the log.",
            "Run /diagnostics. Fix red items before you announce the panel.",
            "Set /permissions roles for staff, admin, and owner.",
            "Set the review channel, log channel, and approval roles.",
            "Read the verification form and change anything that should not be asked.",
            "Post the verification panel in the channel where people should apply.",
            "Set a welcome channel and send yourself a test before you enable it.",
            "Create a backup and store it, the backup password, and the database key away from the server.",
        ],
    )
    add_heading(document, "During the day", 2)
    add_bullets(
        document,
        [
            "Work the queue from the review channel, or glance at counts on the Web UI overview. Actioned today uses the UTC date, not your local midnight.",
            "Use Question when an answer needs a conversation. Use // for notes that must not be sent to the applicant.",
            "Use /modprofile when you want the longer staff record in the log channel.",
            "Cancel a stuck application with /verification cancel-user rather than deleting the review message and hoping the record goes away.",
        ],
    )
    add_heading(document, "Before a change you might want to undo", 2)
    add_bullets(
        document,
        [
            "Create a backup first.",
            "After a restore, restart. The safety copy under data/restore_safety/ is your way back to the previous database file.",
            "If you only needed to undo a form edit, /form reset-verification restores the built-in verification form without a full restore. It does not restore other forms.",
        ],
    )

    add_heading(document, "15. Troubleshooting", 1)
    add_table(
        document,
        ["What you see", "What to check"],
        [
            [
                "The process exits immediately",
                "DISCORD_TOKEN is missing, or TFSBOT_DATABASE_KEY is not 64 hex characters. The log prints the reason. Web UI settings are also checked when WEBUI_ENABLED is true.",
            ],
            [
                "Discord closes the connection on startup",
                "Message Content Intent and Server Members Intent are both enabled in the Developer Portal for this application.",
            ],
            [
                "The bot is online but slash commands are missing",
                "The invite included applications.commands. If TEST_GUILD_ID is set, commands were synced to that server only. A global sync can take a while to appear.",
            ],
            [
                "!ping and custom commands do nothing",
                "Message Content Intent is off, the prefix in BOT_PREFIX is not the one people are typing, or the message is in a questioning thread.",
            ],
            [
                "You need owner permission",
                "Your level is below the command. /permissions my-level shows it. Server owner and BOT_DEV_USER_IDS are owner. Otherwise set the role, or lower that command.",
            ],
            [
                "The bot will not assign or remove a role",
                "The bot role is not above that role, or Manage Roles is missing. /diagnostics names the role.",
            ],
            [
                "Question does not open a thread",
                "In the review channel the bot needs permission to create a public thread and to send messages in it.",
            ],
            [
                "The applicant never receives the DM",
                "They have DMs closed, they do not share a server with the bot any more, or Discord refused the message. The staff action can still have succeeded.",
            ],
            [
                "Welcome never posts",
                "Welcome is disabled, the channel is missing, the bot cannot speak or embed there, or the approval happened before welcome was enabled. Use Test to separate a wording problem from a delivery problem.",
            ],
            [
                "Invite tracking says it is not synchronised",
                "The bot needs Manage Server. The cache is built once after login. A warning here is not the same as a broken verification panel.",
            ],
            [
                "The Web UI does not open",
                "WEBUI_ENABLED is true, the port is free, and you are connecting to the bind address. In Docker, 127.0.0.1 inside the container is not your computer. Compose does not publish the port unless you add that yourself.",
            ],
            [
                "Discord login fails or returns you with an error",
                "Redirect URI, client id, client secret, and WEBUI_DISCORD_GUILD_ID. The person must be in that server and have an allowed role.",
            ],
            [
                "A form submit returns 400",
                "The CSRF token is missing or the page is stale. Reload and try again. Do not bookmark a POST.",
            ],
            [
                "You were sent back to the sign-in page",
                "The session is older than eight hours, or idle for more than an hour, or the cookie was cleared.",
            ],
            [
                "Restore says you need a fresh login",
                "Sign out, sign in, and restore within 10 minutes. Browsing the site does not count as a new login.",
            ],
            [
                "The database will not open",
                "The key does not match the file, or the path in APPLICATION_DB_PATH is wrong. Diagnostics only checks that the file exists and is writable, so use the startup error for the key.",
            ],
            [
                "A backup will not restore",
                "The password is wrong, the file is not a .tfsbackup, or it was damaged. The password is not written down by the bot.",
            ],
            [
                "The site looks wrong after a restore",
                "Restart the process. The running bot is still holding the old database.",
            ],
        ],
    )
    add_paragraph(
        document,
        "Logs are lines of time, level, logger name, and message. "
        "Expected mistakes, such as a missing !ping argument, are not logged as faults. "
        "Unexpected command failures are logged with a traceback. Start there when a button or command replies that something went wrong.",
    )

    add_heading(document, "Keeping this manual", 1)
    add_paragraph(
        document,
        "This file is docs/Sanctuary-Servo-Server-Owner-Manual.docx. "
        "Regenerate it with docs/build_server_owner_manual.py after installing python-docx. "
        "That package is only for rebuilding the manual. The bot does not need it.",
    )
    add_paragraph(
        document,
        "If a later change to the bot disagrees with this manual, trust the running bot and the /diagnostics report, then update the manual.",
    )

    return document


def main() -> None:
    document = build()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
